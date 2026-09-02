from io import BytesIO
from pathlib import Path
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import SecretStr

from kerui_recruit.core.settings import Settings
from kerui_recruit.runtime import create_runtime_app


@pytest.fixture
def client(tmp_path: Path):
    settings = Settings(
        data_root=tmp_path / "data",
        session_token=SecretStr("launch-token"),
    )
    app = create_runtime_app(settings)
    with TestClient(app) as test_client:
        yield test_client


def _headers() -> dict[str, str]:
    return {"X-Kerui-Session": "launch-token"}


def _wait_jd_ready(client: TestClient, revision_id: str) -> None:
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        response = client.get("/api/jd", headers=_headers())
        assert response.status_code == 200
        revision = next((item for item in response.json() if item["revision_id"] == revision_id), None)
        if revision and revision["status"] == "READY":
            assert revision["jd_status"] == "OPEN"
            pending = client.app.state.services.index_sync_service.status()["items"]
            if not any(item["entity_type"] == "jd" and item["entity_id"] == revision["jd_id"] for item in pending):
                return
        time.sleep(.05)
    raise AssertionError("Local JD pipeline did not produce an OPEN, READY revision")


def test_jd_import_and_match_round_trip(client: TestClient) -> None:
    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    )
    assert imported.status_code == 202
    body = imported.json()
    assert body["jd_id"]
    _wait_jd_ready(client, body["revision_id"])

    match = client.post(
        "/api/match/jd",
        json={"revision_id": body["revision_id"], "limit": 20},
        headers=_headers(),
    )
    assert match.status_code == 200


def test_batch_match_multiple_jds(client: TestClient) -> None:
    first = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    ).json()
    second = client.post(
        "/api/jd/import",
        json={"company": "某科技", "title": "算法工程师", "source_text": "Python 算法 大模型"},
        headers=_headers(),
    ).json()
    _wait_jd_ready(client, first["revision_id"])
    _wait_jd_ready(client, second["revision_id"])

    batch = client.post(
        "/api/match/batch",
        json={"revision_ids": [first["revision_id"], second["revision_id"]], "limit": 20},
        headers=_headers(),
    )

    assert batch.status_code == 200
    results = batch.json()["results"]
    assert len(results) == 2
    assert {r["revision_id"] for r in results} == {first["revision_id"], second["revision_id"]}
    assert all(r["run_id"] for r in results)


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_import_jd_from_word_and_excel(client: TestClient) -> None:
    docx = _docx_bytes("Java 后端工程师 3年 Java 本科 金融")
    imported_docx = client.post(
        "/api/jd/import-file",
        files={"file": ("jd.docx", docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"company": "某金融", "title": "Java 后端"},
        headers=_headers(),
    )
    assert imported_docx.status_code == 202
    assert imported_docx.json()["jd_id"]

    xlsx = _xlsx_bytes([["岗位", "算法工程师"], ["要求", "Python 算法 大模型"]])
    imported_xlsx = client.post(
        "/api/jd/import-file",
        files={"file": ("jd.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"company": "某科技", "title": "算法工程师"},
        headers=_headers(),
    )
    assert imported_xlsx.status_code == 202
    assert imported_xlsx.json()["jd_id"]
